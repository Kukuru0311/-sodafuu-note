"""PowerPoint(.pptx)で図形・文字を置くための小さな関数集。

編集用データは .pptx（PowerPoint / Keynote / Googleスライドで開ける）。
ここで作った .pptx を LibreOffice で PDF にし、pdftoppm で PNG にする。
"""
import math
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

FONT = "Zen Maru Gothic"


def no_style(sh):
    """既定のスタイル（影など）を外す。LibreOfficeで影が出るのを防ぐ。"""
    st = sh._element.find(qn("p:style"))
    if st is not None:
        sh._element.remove(st)


def rgb(hexstr):
    hexstr = hexstr.lstrip("#")
    return RGBColor(int(hexstr[0:2], 16), int(hexstr[2:4], 16), int(hexstr[4:6], 16))


class Doc:
    """unit: 1単位あたりのEMU。px(96dpi)なら9525、mmなら36000。"""

    def __init__(self, width, height, unit):
        self.u = unit
        self.prs = Presentation()
        self.prs.slide_width = Emu(int(width * unit))
        self.prs.slide_height = Emu(int(height * unit))
        self.W, self.H = width, height

    def e(self, v):
        return Emu(int(v * self.u))

    def slide(self, bg=None, name=None):
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        if bg:
            fill = s.background.fill
            fill.solid()
            fill.fore_color.rgb = rgb(bg)
        if name:
            s.notes_slide.notes_text_frame.text = name
        return s

    def shape(self, s, kind, x, y, w, h, fill=None, line=None, line_w=0, radius=None, name=None, rot=0):
        sh = s.shapes.add_shape(kind, self.e(x), self.e(y), self.e(w), self.e(h))
        if fill:
            sh.fill.solid()
            sh.fill.fore_color.rgb = rgb(fill)
        else:
            sh.fill.background()
        if line:
            sh.line.color.rgb = rgb(line)
            sh.line.width = self.e(line_w)
        else:
            sh.line.fill.background()
        no_style(sh)
        if radius is not None and kind == MSO_SHAPE.ROUNDED_RECTANGLE:
            sh.adjustments[0] = radius
        if name:
            sh.name = name
        if rot:
            sh.rotation = rot
        sh.text_frame.text = ""
        return sh

    def rect(self, s, x, y, w, h, fill, **kw):
        return self.shape(s, MSO_SHAPE.RECTANGLE, x, y, w, h, fill, **kw)

    def rrect(self, s, x, y, w, h, fill, radius=0.15, **kw):
        return self.shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill, radius=radius, **kw)

    def circle(self, s, cx, cy, r, fill, **kw):
        return self.shape(s, MSO_SHAPE.OVAL, cx - r, cy - r, 2 * r, 2 * r, fill, **kw)

    def wave(self, s, x, y, w, h, amp, waves, fill, top=True, name=None):
        """上辺（top=True）が波になった帯。"""
        n = 80
        pts = []
        for i in range(n + 1):
            t = i / n
            px = x + w * t
            py = y + amp * math.sin(t * waves * 2 * math.pi)
            pts.append((px, py))
        if top:
            pts += [(x + w, y + h), (x, y + h)]
        fb = s.shapes.build_freeform(self.e(pts[0][0]), self.e(pts[0][1]), scale=1.0)
        fb.add_line_segments([(self.e(a), self.e(b)) for a, b in pts[1:]], close=True)
        sh = fb.convert_to_shape()
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb(fill)
        sh.line.fill.background()
        no_style(sh)
        if name:
            sh.name = name
        return sh

    def scribble(self, s, x, y, w, amp, loops, color, width, name=None):
        """手描き風のくるくる線（下線・飾り用）。"""
        n = 160
        pts = []
        for i in range(n + 1):
            t = i / n
            ang = t * loops * 2 * math.pi
            px = x + w * t + amp * 0.9 * math.cos(ang)
            py = y + amp * math.sin(ang)
            pts.append((px, py))
        fb = s.shapes.build_freeform(self.e(pts[0][0]), self.e(pts[0][1]), scale=1.0)
        fb.add_line_segments([(self.e(a), self.e(b)) for a, b in pts[1:]], close=False)
        sh = fb.convert_to_shape()
        sh.fill.background()
        sh.line.color.rgb = rgb(color)
        sh.line.width = self.e(width)
        no_style(sh)
        if name:
            sh.name = name
        return sh

    def text(self, s, x, y, w, h, content, size, color="#333333", bold=False,
             align="left", anchor="top", spacing=1.25, name=None, fill=None):
        """content: 文字列 or [(文字列, {size,bold,color})...] の段落リスト。size は pt。"""
        tb = s.shapes.add_textbox(self.e(x), self.e(y), self.e(w), self.e(h))
        if fill:
            tb.fill.solid()
            tb.fill.fore_color.rgb = rgb(fill)
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE,
                              "bottom": MSO_ANCHOR.BOTTOM}[anchor]
        paras = content if isinstance(content, list) else [(line, {}) for line in content.split("\n")]
        for i, item in enumerate(paras):
            if isinstance(item, str):
                item = (item, {})
            t, opt = item
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER,
                           "right": PP_ALIGN.RIGHT}[opt.get("align", align)]
            p.line_spacing = opt.get("spacing", spacing)
            if opt.get("space_before"):
                p.space_before = Pt(opt["space_before"])
            r = p.add_run()
            r.text = t
            f = r.font
            f.size = Pt(opt.get("size", size))
            f.bold = opt.get("bold", bold)
            f.color.rgb = rgb(opt.get("color", color))
            f.name = FONT
            rPr = r._r.get_or_add_rPr()
            for tag in ("a:ea", "a:cs"):
                el = rPr.makeelement(qn(tag), {"typeface": FONT})
                rPr.append(el)
        if name:
            tb.name = name
        return tb

    def save(self, path):
        self.prs.save(path)
